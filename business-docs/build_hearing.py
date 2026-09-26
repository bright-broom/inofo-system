"""ヒアリングシート（無料相談・かんたん診断用）の生成スクリプト。

.venv/bin/python build_hearing.py で out/ヒアリングシート.xlsx を再生成する。
「ヒアリング」で基本情報とお困りごとを聞き、「かんたん診断」のチェックに答えると、
要対応の件数と推奨プランが自動で出る。結果は提案書の2章・3章にそのまま転記できる。
記入例は「サンプル商事株式会社（従業員25名）」で、提案書・月次報告書・請求書の例と同じ。
"""
from datetime import datetime

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

from pricing import PLANS

FONT = "游ゴシック"
INPUT_FILL = PatternFill("solid", fgColor="FFF59D")
HEAD_FILL = PatternFill("solid", fgColor="1F1F1F")
SEC_FILL = PatternFill("solid", fgColor="E8E6DC")
SOFT_FILL = PatternFill("solid", fgColor="F2F2F2")
ALERT_FILL = PatternFill("solid", fgColor="FBD5D2")
THIN = Side(style="thin", color="999999")
BOX = Border(top=THIN, bottom=THIN, left=THIN, right=THIN)

S = "'設定'!"
H = "'ヒアリング'!"
C = "'かんたん診断'!"


def put(ws, ref, value, *, size=10, bold=False, fmt=None, align=None, fill=None, border=None, color="111111", wrap=False):
    c = ws[ref]
    c.value = value
    c.font = Font(name=FONT, size=size, bold=bold, color=color)
    if fmt:
        c.number_format = fmt
    c.alignment = Alignment(horizontal=align, vertical="center", wrap_text=wrap)
    if fill:
        c.fill = fill
    if border:
        c.border = border
    return c


def inp(ws, ref, value, **kw):
    return put(ws, ref, value, fill=INPUT_FILL, color="0000CC", **kw)


def print_setup(ws, last_col, last_row):
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.orientation = "portrait"
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.print_area = f"A1:{last_col}{last_row}"
    ws.sheet_view.showGridLines = False


# ---------------------------------------------------------------- 設定
def build_settings(wb):
    ws = wb.create_sheet("設定")
    ws.sheet_view.showGridLines = False
    for col, w in {"A": 2, "B": 22, "C": 16, "D": 16, "E": 16}.items():
        ws.column_dimensions[col].width = w
    put(ws, "B2", "推奨プランの判定基準", size=14, bold=True)
    put(ws, "B3", "従業員数と、月のIT対応時間の推定の「大きい方」でプランを選ぶ。値はサイト・契約書と同じ。", size=9, color="666666")

    for col, lab in zip("BCDE", ["プラン", "従業員数の上限", "月間対応時間", "月額（税別）"]):
        put(ws, f"{col}5", lab, bold=True, fill=SOFT_FILL, border=BOX)
    for i, (name, emp, hours, fee) in enumerate([
        *[(p["name"], p["employees"][1], p["hours"], p["price"]) for p in PLANS], ("要相談", None, None, None),
    ]):
        r = 6 + i
        put(ws, f"B{r}", name, border=BOX)
        if emp is not None:
            inp(ws, f"C{r}", emp, fmt='0"名"', border=BOX)
            inp(ws, f"D{r}", hours, fmt='0"時間"', border=BOX)
            inp(ws, f"E{r}", fee, fmt='"¥"#,##0', border=BOX)
        else:
            for col in "CDE":
                put(ws, f"{col}{r}", "—", border=BOX, align="center", color="888888")

    # 判定の途中計算（ヒアリングシートから参照）
    put(ws, "B11", "判定の途中計算", size=12, bold=True)
    put(ws, "B12", "従業員数からの段階", border=BOX, fill=SOFT_FILL)
    put(ws, "C12", f'=IF({H}$C$9="","",IF({H}$C$9<=$C$6,1,IF({H}$C$9<=$C$7,2,IF({H}$C$9<=$C$8,3,4))))', border=BOX, align="center")
    put(ws, "B13", "対応時間からの段階", border=BOX, fill=SOFT_FILL)
    put(ws, "C13", f'=IF({H}$C$16="","",IF({H}$C$16<=$D$6,1,IF({H}$C$16<=$D$7,2,IF({H}$C$16<=$D$8,3,4))))', border=BOX, align="center")
    put(ws, "B14", "推奨の段階", border=BOX, fill=SOFT_FILL, bold=True)
    put(ws, "C14", '=IF(AND(C12="",C13=""),"",MAX(C12,C13))', border=BOX, align="center", bold=True)

    put(ws, "B17", "選択肢", size=12, bold=True)
    lists = {
        "B": ["はい", "いいえ", "不明"],
        "C": ["兼任", "専任", "いない"],
        "D": ["Microsoft 365", "Google Workspace", "その他", "なし"],
        "E": ["オンライン", "訪問", "電話"],
    }
    for col, items in lists.items():
        for i, v in enumerate(items):
            put(ws, f"{col}{18 + i}", v, border=BOX)
    return ws


# ---------------------------------------------------------------- ヒアリング
def build_hearing(wb):
    ws = wb.active
    ws.title = "ヒアリング"
    for col, w in {"A": 2, "B": 24, "C": 22, "D": 20, "E": 30}.items():
        ws.column_dimensions[col].width = w

    dv_yn = DataValidation(type="list", formula1=f"={S}$B$18:$B$20", allow_blank=True)
    dv_role = DataValidation(type="list", formula1=f"={S}$C$18:$C$20", allow_blank=True)
    dv_suite = DataValidation(type="list", formula1=f"={S}$D$18:$D$21", allow_blank=True)
    dv_how = DataValidation(type="list", formula1=f"={S}$E$18:$E$20", allow_blank=True)
    dv_num = DataValidation(type="decimal", operator="between", formula1="0", formula2="10000", allow_blank=True,
                            error="数値を入力してください", showErrorMessage=True)
    for dv in (dv_yn, dv_role, dv_suite, dv_how, dv_num):
        ws.add_data_validation(dv)

    put(ws, "B1", "ヒアリングシート", size=18, bold=True)
    put(ws, "B2", "黄色のセルに記入します。▼はプルダウンです。社外に出さない社内用メモです。", size=8, color="999999")

    def section(r, title):
        put(ws, f"B{r}", title, size=11, bold=True, fill=SEC_FILL)
        for col in "CDE":
            ws[f"{col}{r}"].fill = SEC_FILL

    def row(r, label, value=None, *, dv=None, fmt=None, hint=None, span=True):
        put(ws, f"B{r}", label, bold=True, size=9, border=BOX, fill=SOFT_FILL)
        inp(ws, f"C{r}", value, fmt=fmt, size=10, border=BOX, wrap=True)
        if dv:
            dv.add(f"C{r}")
        if hint:
            put(ws, f"E{r}", hint, size=8, color="777777", wrap=True)
            if span:
                ws.merge_cells(f"C{r}:D{r}")
        elif span:
            ws.merge_cells(f"C{r}:E{r}")

    # 1. 基本情報
    section(4, "1. 基本情報")
    row(5, "会社名", "サンプル商事株式会社")
    row(6, "ご担当者（部署・お名前）", "総務部 A様")
    row(7, "決裁者", "代表 C様", hint="提案書の宛先。同席されたか")
    row(8, "業種", "卸売業")
    row(9, "従業員数（名）", 25, dv=dv_num, fmt='0"名"', hint="推奨プランの判定に使う")
    row(10, "拠点数・所在地", "本社1拠点（東京都〇〇区）")
    row(11, "実施日時・方法", "2026/9/28 19:30〜 オンライン", hint="オンライン／訪問／電話")

    # 2. IT対応の体制
    section(13, "2. IT対応の体制")
    row(14, "IT担当者の有無", "兼任", dv=dv_role, hint="兼任／専任／いない")
    row(15, "担当者の本来の業務", "総務・経理")
    row(16, "月のIT対応時間の推定（時間）", 12, dv=dv_num, fmt='0"時間"', hint="担当者に「月に何時間くらい？」と聞く。推奨プランの判定に使う")
    row(17, "付き合いのあるIT業者", "複合機の保守会社のみ")

    # 3. IT環境
    section(19, "3. IT環境")
    row(20, "PCの台数・OS", "25台（Windows 11 が20台、Windows 10 が5台）")
    row(21, "メール・グループウェア", "Microsoft 365", dv=dv_suite, hint="Microsoft 365／Google Workspace／その他／なし")
    row(22, "ファイルの保存場所", "Microsoft 365 と、個人契約のファイル共有が混在")
    row(23, "主に使っているSaaS", "会計ソフト、勤怠管理、チャット")
    row(24, "ネットワーク・Wi-Fi", "Wi-Fi 1系統（来客と共用）")
    row(25, "その他の機器", "複合機1台、NAS 1台")

    # 4. お困りごと
    section(27, "4. お困りごと（困っている順に）")
    row(28, "1番目", "退職者のアカウントが残っている／入退社のたびに手間がかかる")
    row(29, "2番目", "取引先からセキュリティ確認書が届いたが答えられない")
    row(30, "3番目", "担当者（A様）が休むとITのことが誰も分からない")
    row(31, "直近で起きたトラブル", "共有フォルダにアクセスできなくなり、半日業務が止まった")

    # 5. 今後の予定・ご希望
    section(33, "5. 今後の予定・ご希望")
    row(34, "採用・移転などの予定", "来年4月に3名採用予定")
    row(35, "導入を考えているツール", "特になし")
    row(36, "希望する開始時期", "2026年11月")
    row(37, "予算感", "月5〜7万円程度")
    row(38, "その他・メモ", "代表は情報漏えいを特に気にしている")

    # 6. 判定（自動）
    section(40, "6. 判定（自動）")
    put(ws, "B41", "かんたん診断：要対応", bold=True, size=9, border=BOX, fill=SOFT_FILL)
    put(ws, "C41", f'="重要度・高 "&{C}$G$4&"件／中 "&{C}$G$5&"件"', size=10, bold=True, border=BOX)
    ws.merge_cells("C41:E41")
    put(ws, "B42", "推奨プラン", bold=True, size=9, border=BOX, fill=SOFT_FILL)
    put(ws, "C42", f'=IF({S}$C$14="","（従業員数か対応時間を入力）",INDEX({S}$B$6:$B$9,{S}$C$14))', size=12, bold=True, border=BOX, color="3D38E0")
    put(ws, "D42", f'=IF(OR({S}$C$14="",{S}$C$14=4),"","月額 "&TEXT(INDEX({S}$E$6:$E$8,{S}$C$14),"#,##0")&"円／月"&INDEX({S}$D$6:$D$8,{S}$C$14)&"時間")', size=9, border=BOX)
    ws.merge_cells("D42:E42")
    put(ws, "B43", "判定メモ", bold=True, size=9, border=BOX, fill=SOFT_FILL)
    put(
        ws, "C43",
        f'=IF({S}$C$14="","",IF({S}$C$14=4,"規模・作業量が上限を超えます。ひとり運営で受けられるか慎重に判断。",'
        f'IF({S}$C$13>{S}$C$12,"対応時間の推定から、従業員数の目安より上のプランを推奨。",'
        f'"従業員数と対応時間の推定の両方から判定。"))&IF({C}$G$4>=3," 重要度・高の要対応が多いため、最初の3ヶ月は土台づくりを優先。",""))',
        size=9, border=BOX, wrap=True,
    )
    ws.merge_cells("C43:E43")
    ws.row_dimensions[43].height = 32

    put(ws, "B45", "→ 結果は提案書の「2. 現状と課題」「3. ご提案内容」に転記します。", size=8, color="666666")
    print_setup(ws, "E", 45)
    return ws


# ---------------------------------------------------------------- かんたん診断
CHECKS = [
    # (区分, 質問, 重要度, 記入例の回答, 記入例のメモ)
    ("アカウント", "退職者のアカウントは、退職日に停止していますか", "高", "いいえ", "2件残っている"),
    ("アカウント", "管理者アカウントは、人ごとに分かれていますか（共有していない）", "中", "いいえ", "A様と代表で共有"),
    ("アカウント", "入社・退職時のIT手続きが、手順書になっていますか", "中", "いいえ", ""),
    ("認証", "メールやクラウドに、多要素認証を導入していますか", "高", "いいえ", ""),
    ("認証", "パスワードの使い回しを禁止し、守られていますか", "中", "不明", ""),
    ("PC", "すべてのPCが、サポート期間内のOSで最新の状態ですか", "高", "いいえ", "Windows 10 が5台"),
    ("PC", "すべてのPCに、ウイルス対策（EDRなど）が入っていますか", "高", "はい", "OS標準の機能"),
    ("PC", "誰がどのPCを使っているか、台帳がありますか", "中", "いいえ", ""),
    ("データ", "重要なデータのバックアップを取っていますか", "高", "はい", "NASに週1回"),
    ("データ", "バックアップから実際に復元できることを確認しましたか", "中", "不明", ""),
    ("データ", "業務のファイルは、会社で契約したサービスだけで扱っていますか", "中", "いいえ", "個人契約のファイル共有あり"),
    ("データ", "共有フォルダのアクセス権は、部署ごとに整理されていますか", "中", "いいえ", ""),
    ("ネットワーク", "Wi-Fiは、社内用と来客用に分かれていますか", "中", "いいえ", ""),
    ("運用", "不審なメールを受け取ったときの連絡先が決まっていますか", "中", "いいえ", ""),
    ("運用", "契約中のソフト・SaaSの一覧（契約者・費用）がありますか", "中", "はい", ""),
    ("運用", "取引先からのセキュリティ確認書に回答できますか", "中", "いいえ", "9月に届いたものが未回答"),
]
CHK_FIRST = 8


def build_check(wb):
    ws = wb.create_sheet("かんたん診断")
    for col, w in {"A": 2, "B": 5, "C": 12, "D": 54, "E": 8, "F": 9, "G": 10, "H": 26}.items():
        ws.column_dimensions[col].width = w
    put(ws, "B1", "かんたん診断チェックリスト", size=16, bold=True)
    put(ws, "B2", "「はい」が望ましい状態です。「いいえ」「不明」は要対応として数えます。", size=8, color="999999")

    last = CHK_FIRST + len(CHECKS) - 1
    rng_ans, rng_lv, rng_judge = f"$F${CHK_FIRST}:$F${last}", f"$E${CHK_FIRST}:$E${last}", f"$G${CHK_FIRST}:$G${last}"
    put(ws, "D4", "要対応（重要度・高）", bold=True, size=9, align="right")
    put(ws, "G4", f'=COUNTIFS({rng_lv},"高",{rng_judge},"要対応")', fmt='0"件"', bold=True, border=BOX, align="center", fill=ALERT_FILL)
    put(ws, "D5", "要対応（重要度・中）", bold=True, size=9, align="right")
    put(ws, "G5", f'=COUNTIFS({rng_lv},"中",{rng_judge},"要対応")', fmt='0"件"', bold=True, border=BOX, align="center")
    put(ws, "D6", "未回答", size=9, align="right", color="666666")
    put(ws, "G6", f'=COUNTBLANK({rng_ans})', fmt='0"件"', size=9, border=BOX, align="center")

    hdr = CHK_FIRST - 1
    for col, lab in zip("BCDEFGH", ["No", "区分", "確認すること", "重要度", "回答", "判定", "メモ"]):
        put(ws, f"{col}{hdr}", lab, bold=True, color="FFFFFF", fill=HEAD_FILL, align="center", border=BOX, size=9)

    dv = DataValidation(type="list", formula1=f"={S}$B$18:$B$20", allow_blank=True)
    ws.add_data_validation(dv)
    for i, (cat, q, lv, ans, memo) in enumerate(CHECKS):
        r = CHK_FIRST + i
        put(ws, f"B{r}", i + 1, size=9, border=BOX, align="center")
        put(ws, f"C{r}", cat, size=9, border=BOX)
        put(ws, f"D{r}", q, size=9, border=BOX, wrap=True)
        put(ws, f"E{r}", lv, size=9, border=BOX, align="center", bold=lv == "高")
        inp(ws, f"F{r}", ans, size=9, border=BOX, align="center")
        dv.add(f"F{r}")
        put(ws, f"G{r}", f'=IF(F{r}="","",IF(F{r}="はい","OK","要対応"))', size=9, border=BOX, align="center", bold=True)
        inp(ws, f"H{r}", memo or None, size=9, border=BOX, wrap=True)
        ws.row_dimensions[r].height = 26
    ws.conditional_formatting.add(f"G{CHK_FIRST}:G{last}", FormulaRule(formula=[f'G{CHK_FIRST}="要対応"'], fill=ALERT_FILL))

    put(ws, f"B{last + 2}", "→ 要対応の項目（特に重要度・高）を、提案書の「2. 現状と課題」と「4. 最初の3ヶ月の進め方」に反映します。", size=8, color="666666")
    print_setup(ws, "H", last + 2)
    return ws


def main():
    wb = Workbook()
    build_hearing(wb)
    build_check(wb)
    build_settings(wb)
    wb.calculation.fullCalcOnLoad = True
    out = "out/ヒアリングシート.xlsx"
    wb.save(out)
    print(out)


if __name__ == "__main__":
    main()

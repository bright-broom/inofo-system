"""オンボーディング（契約〜運用開始〜最初の1ヶ月）チェックリストの生成スクリプト。

.venv/bin/python build_onboarding.py で out/オンボーディングチェックリスト.xlsx を再生成する。
開始日を入れると各タスクの期限が自動で決まり、担当ごとの進み具合と期限切れを表示する。
記入例は「サンプル商事株式会社（スタンダード・11/1開始）」で、他のテンプレートの例と同じ。
"""
from datetime import date

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

FONT = "游ゴシック"
INPUT_FILL = PatternFill("solid", fgColor="FFF59D")
HEAD_FILL = PatternFill("solid", fgColor="1F1F1F")
SEC_FILL = PatternFill("solid", fgColor="E8E6DC")
SOFT_FILL = PatternFill("solid", fgColor="F2F2F2")
ALERT_FILL = PatternFill("solid", fgColor="FBD5D2")
DONE_FONT = Font(name=FONT, size=9, color="9A9A9A")
THIN = Side(style="thin", color="999999")
BOX = Border(top=THIN, bottom=THIN, left=THIN, right=THIN)

MD = 'm"/"d"("aaa")"'
STATUSES = ["未着手", "進行中", "完了", "不要"]
OWNERS = ["ラクシス", "お客様", "両者"]

# (フェーズ, タスク, 担当, 期限の決め方, 根拠・メモ, 記入例の状態)
# 期限の決め方：整数 = 基準日からの日数（マイナスは前）／"M5" = 基準月の翌月5日／"P" = 基準月の翌月末
TASKS = [
    ("1. 契約", "秘密保持契約を締結する", "両者", -14, "管理者権限を預かる前に必ず先に結ぶ", "完了"),
    ("1. 契約", "業務委託契約を締結する（別紙1のプラン・開始日を確定）", "両者", -10, "契約書 第16条：初月は日割り", "完了"),
    ("1. 契約", "請求書の宛名・送付先・ご担当者を確認する", "ラクシス", -10, "インボイス登録番号の要否もここで確認", "完了"),
    ("1. 契約", "社内の窓口担当者を決める", "お客様", -10, "ラクシスへの依頼をまとめる方", "完了"),
    ("2. 準備", "連絡手段を決め、チャンネルを用意する（メール／チャット）", "両者", -7, "契約書 別紙1「連絡手段」", "進行中"),
    ("2. 準備", "ラクシス専用の管理者アカウントを発行する（共有しない）", "お客様", -5, "契約書 第6条。既存の管理者アカウントの共有は避ける", "未着手"),
    ("2. 準備", "認証情報の受け渡し方法を決める（メール本文に書かない）", "両者", -5, "パスワード管理ツールの共有機能などを使う", "未着手"),
    ("2. 準備", "データのバックアップ状況を確認する", "両者", -5, "契約書 第5条：バックアップはお客様の責任", "未着手"),
    ("2. 準備", "付き合いのある業者の連絡先を一覧にする（複合機保守など）", "お客様", -3, "範囲外の作業を業者へ回すときに使う", "未着手"),
    ("2. 準備", "社内向けの案内文を送る", "ラクシス", -3, "営業メールテンプレート 7「運用開始のご案内」", "未着手"),
    ("2. 準備", "不在予定と、対応時間外の扱いを説明する", "ラクシス", -3, "契約書 第4条：即時対応・24時間監視は範囲外", "未着手"),
    ("3. 開始", "社内に窓口と依頼方法をアナウンスする", "お客様", 0, "案内文をそのまま転送でよい", "未着手"),
    ("3. 開始", "キックオフ（オンライン30分）を行う", "両者", 0, "ヒアリングシートの要対応項目を共有", "未着手"),
    ("3. 開始", "受け取った管理者権限でログインできることを確認する", "ラクシス", 0, "", "未着手"),
    ("4. 最初の1ヶ月", "アカウントを棚卸しする（退職者・不要アカウントの洗い出し）", "ラクシス", 7, "提案書「最初の3ヶ月の進め方」1ヶ月目", "未着手"),
    ("4. 最初の1ヶ月", "PCとIT資産の台帳を作る", "ラクシス", 14, "ヒアリングシートの「IT環境」を出発点にする", "未着手"),
    ("4. 最初の1ヶ月", "契約中のソフト・SaaSの一覧を作る", "両者", 14, "契約者と費用が分かるように", "未着手"),
    ("4. 最初の1ヶ月", "重要度・高の要対応項目の着手順を決める", "両者", 14, "かんたん診断の結果から", "未着手"),
    ("4. 最初の1ヶ月", "手順書・台帳の保存場所を決める", "両者", 14, "成果物はお客様の環境に置く（契約書 第11条）", "未着手"),
    ("4. 最初の1ヶ月", "初回の請求書を送る", "ラクシス", "M5", "契約書 第8条：翌月5日まで", "未着手"),
    ("4. 最初の1ヶ月", "初回の月次レポートを送る", "ラクシス", "M10", "契約書 第9条：翌月10日まで", "未着手"),
    ("4. 最初の1ヶ月", "1ヶ月の振り返り（定例）で、プランの過不足を確認する", "両者", 30, "作業時間が枠を超えていれば上位プランを相談", "未着手"),
    ("5. ラクシス社内", "顧客フォルダを作り、契約書類を保存する", "ラクシス", -10, "", "完了"),
    ("5. ラクシス社内", "月次報告書・請求書のテンプレートを顧客用に複製する", "ラクシス", -3, "", "未着手"),
    ("5. ラクシス社内", "預かった認証情報をパスワード管理ツールに登録する", "ラクシス", 0, "暗号化して保管。PCのメモ帳などに残さない", "未着手"),
]
FIRST = 10


def put(ws, ref, value, *, size=10, bold=False, fmt=None, align=None, fill=None, border=None, color="111111", wrap=False):
    c = ws[ref]
    c.value = value
    c.font = Font(name=FONT, size=size, bold=bold, color=color)
    if fmt:
        c.number_format = fmt
    c.alignment = Alignment(horizontal=align, vertical="center", wrap_text=wrap)
    if fill:
        c.fill = fill
    if border:
        c.border = border
    return c


def inp(ws, ref, value, **kw):
    return put(ws, ref, value, fill=INPUT_FILL, color="0000CC", **kw)


def build_checklist(title, note, info_rows, tasks, out):
    """info_rows の3行目（D5）を基準日とし、各タスクの期限を計算する。"""
    wb = Workbook()
    ws = wb.active
    ws.title = "チェックリスト"
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = f"A{FIRST}"
    for col, w in {"A": 2, "B": 4, "C": 13, "D": 50, "E": 10, "F": 11, "G": 9, "H": 10, "I": 36}.items():
        ws.column_dimensions[col].width = w

    last = FIRST + len(tasks) - 1
    rng_owner, rng_st, rng_flag = f"$E${FIRST}:$E${last}", f"$G${FIRST}:$G${last}", f"$H${FIRST}:$H${last}"

    put(ws, "B1", title, size=16, bold=True)
    put(ws, "B2", note, size=8, color="999999")

    # 基本情報
    for r, (lab, val, fmt) in enumerate(info_rows, start=3):
        put(ws, f"C{r}", lab, bold=True, size=9, border=BOX, fill=SOFT_FILL)
        inp(ws, f"D{r}", val, fmt=fmt, border=BOX)

    # 進み具合
    put(ws, "F3", "進み具合", bold=True, size=9, border=BOX, fill=SOFT_FILL, align="center")
    ws.merge_cells("F3:G3")
    put(ws, "H3", f'=IF(COUNTA({rng_st})-COUNTIF({rng_st},"不要")=0,0,COUNTIF({rng_st},"完了")/(COUNTA({rng_st})-COUNTIF({rng_st},"不要")))',
        fmt="0%", bold=True, size=12, border=BOX, align="center")
    for i, who in enumerate(OWNERS):
        r = 4 + i
        put(ws, f"F{r}", f"{who}の残り", size=9, border=BOX, fill=SOFT_FILL, align="center")
        ws.merge_cells(f"F{r}:G{r}")
        put(ws, f"H{r}", f'=COUNTIFS({rng_owner},"{who}",{rng_st},"<>完了",{rng_st},"<>不要")', fmt='0"件"', size=10, border=BOX, align="center")
    put(ws, "F7", "期限切れ", size=9, bold=True, border=BOX, fill=ALERT_FILL, align="center")
    ws.merge_cells("F7:G7")
    put(ws, "H7", f'=COUNTIF({rng_flag},"期限切れ")', fmt='0"件"', bold=True, size=10, border=BOX, align="center")

    # 一覧
    hdr = FIRST - 1
    for col, lab in zip("BCDEFGHI", ["No", "フェーズ", "やること", "担当", "期限", "状態", "注意", "根拠・メモ"]):
        put(ws, f"{col}{hdr}", lab, bold=True, color="FFFFFF", fill=HEAD_FILL, align="center", border=BOX, size=9)

    dv_st = DataValidation(type="list", formula1='"' + ",".join(STATUSES) + '"', allow_blank=True)
    dv_ow = DataValidation(type="list", formula1='"' + ",".join(OWNERS) + '"', allow_blank=True)
    ws.add_data_validation(dv_st)
    ws.add_data_validation(dv_ow)

    for i, (phase, task, owner, due, memo, status) in enumerate(tasks):
        r = FIRST + i
        put(ws, f"B{r}", i + 1, size=9, border=BOX, align="center")
        put(ws, f"C{r}", phase, size=9, border=BOX)
        put(ws, f"D{r}", task, size=9, border=BOX, wrap=True)
        inp(ws, f"E{r}", owner, size=9, border=BOX, align="center")
        dv_ow.add(f"E{r}")
        if due == "P":  # 基準月の翌月末
            fx = '=IF($D$5="","",EOMONTH($D$5,1))'
        elif isinstance(due, str):  # "M5" = 基準月の翌月5日
            fx = f'=IF($D$5="","",EOMONTH($D$5,0)+{int(due[1:])})'
        else:
            fx = f'=IF($D$5="","",$D$5{due:+d})' if due else '=IF($D$5="","",$D$5)'
        put(ws, f"F{r}", fx, fmt=MD, size=9, border=BOX, align="center")
        inp(ws, f"G{r}", status, size=9, border=BOX, align="center")
        dv_st.add(f"G{r}")
        put(ws, f"H{r}", f'=IF(OR(F{r}="",G{r}="完了",G{r}="不要"),"",IF(F{r}<TODAY(),"期限切れ",IF(F{r}-TODAY()<=3,"まもなく","")))',
            size=9, border=BOX, align="center", bold=True)
        put(ws, f"I{r}", memo or None, size=8, border=BOX, wrap=True, color="555555")
        ws.row_dimensions[r].height = 28

    # 期限切れは赤、まもなくは黄、完了・不要の行は薄く
    ws.conditional_formatting.add(f"H{FIRST}:H{last}", FormulaRule(formula=[f'H{FIRST}="期限切れ"'], fill=ALERT_FILL))
    ws.conditional_formatting.add(f"H{FIRST}:H{last}", FormulaRule(formula=[f'H{FIRST}="まもなく"'], fill=INPUT_FILL))
    ws.conditional_formatting.add(f"B{FIRST}:F{last}", FormulaRule(formula=[f'OR($G{FIRST}="完了",$G{FIRST}="不要")'], font=DONE_FONT))

    put(ws, f"B{last + 2}", "※ 期限は目安です。状況に合わせて「状態」を「不要」にしたり、行を追加したりしてください。", size=8, color="666666")

    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.orientation = "landscape"
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.print_area = f"A1:I{last + 2}"
    ws.print_title_rows = f"{hdr}:{hdr}"

    wb.calculation.fullCalcOnLoad = True
    wb.save(out)
    print(out)


DATE_FMT = 'yyyy"年"m"月"d"日("aaa")"'


def main():
    build_checklist(
        "オンボーディング チェックリスト",
        "黄色のセルに記入します。開始日を入れると期限が自動で決まります。お客様と共有するときはPDFにして送ります。",
        [
            ("お客様", "サンプル商事株式会社", None),
            ("プラン", "スタンダード", None),
            ("開始日", date(2026, 11, 1), DATE_FMT),
            ("お客様の窓口", "総務部 A様", None),
            ("ラクシス担当", "〔氏名〕", None),
        ],
        TASKS,
        "out/オンボーディングチェックリスト.xlsx",
    )


if __name__ == "__main__":
    main()

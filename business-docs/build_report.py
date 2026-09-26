"""月次作業報告書テンプレート（Excel）の生成スクリプト。

.venv/bin/python build_report.py で out/月次作業報告書.xlsx を再生成する。
「作業記録」シートに日々の作業を書くと、「月次レポート」シートに自動集計される。
契約書 第9条（対応件数・作業時間・主な作業内容を翌月10日までに報告）に対応。
"""
from datetime import date

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

from pricing import OVERAGE
from pricing import PLANS as _PLANS

FONT = "游ゴシック"
INPUT_FILL = PatternFill("solid", fgColor="FFF59D")
HEAD_FILL = PatternFill("solid", fgColor="1F1F1F")
SOFT_FILL = PatternFill("solid", fgColor="F2F2F2")
THIN = Side(style="thin", color="999999")
BOX = Border(top=THIN, bottom=THIN, left=THIN, right=THIN)
BOTTOM = Border(bottom=Side(style="medium", color="111111"))

YEN = '"¥"#,##0'
HOURS = '0.00"h"'
DATE = 'yyyy"年"m"月"d"日"'
MD = 'm"/"d'
MONTH = 'yyyy"年"m"月"'

LOG_FIRST, LOG_LAST = 5, 64  # 作業記録の行（60件分）
CATEGORIES = ["ヘルプデスク", "アカウント管理", "PCセットアップ", "クラウド運用", "セキュリティ", "ネットワーク", "業務自動化", "定例・相談", "その他"]
STATUSES = ["完了", "対応中", "保留"]
PLANS = [(p["name"], p["hours"], p["price"]) for p in _PLANS]

L = "'作業記録'!"
S = "'設定'!"


def put(ws, ref, value, *, size=10, bold=False, fmt=None, align=None, fill=None, border=None, color="111111", wrap=False):
    c = ws[ref]
    c.value = value
    c.font = Font(name=FONT, size=size, bold=bold, color=color)
    if fmt:
        c.number_format = fmt
    c.alignment = Alignment(horizontal=align, vertical="center", wrap_text=wrap)
    if fill:
        c.fill = fill
    if border:
        c.border = border
    return c


def inp(ws, ref, value, **kw):
    return put(ws, ref, value, fill=INPUT_FILL, color="0000CC", **kw)


# ---------------------------------------------------------------- 設定
def build_settings(wb):
    ws = wb.create_sheet("設定")
    ws.sheet_view.showGridLines = False
    for col, w in {"A": 2, "B": 18, "C": 16, "D": 16, "E": 4, "F": 18}.items():
        ws.column_dimensions[col].width = w
    put(ws, "B2", "プランと区分の設定", size=14, bold=True)
    put(ws, "B3", "料金は契約書 別紙1と同じ。変更したら契約書・見積書側もそろえること。", size=9, color="666666")

    for col, label in zip("BCD", ["プラン", "月間対応時間", "月額（税別）"]):
        put(ws, f"{col}5", label, bold=True, fill=SOFT_FILL, border=BOX)
    for i, (name, hours, fee) in enumerate(PLANS):
        r = 6 + i
        put(ws, f"B{r}", name, border=BOX)
        inp(ws, f"C{r}", hours, fmt='0"時間"', border=BOX)
        inp(ws, f"D{r}", fee, fmt=YEN, border=BOX)
    put(ws, "B10", "超過単価（税別）", bold=True, fill=SOFT_FILL, border=BOX)
    inp(ws, "C10", OVERAGE, fmt=YEN, border=BOX)

    put(ws, "F5", "作業の区分", bold=True, fill=SOFT_FILL, border=BOX)
    for i, c in enumerate(CATEGORIES):
        put(ws, f"F{6 + i}", c, border=BOX)
    return ws


# ---------------------------------------------------------------- 作業記録
def build_log(wb):
    ws = wb.create_sheet("作業記録")
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = f"A{LOG_FIRST}"
    for col, w in {"A": 5, "B": 9, "C": 16, "D": 15, "E": 46, "F": 10, "G": 9}.items():
        ws.column_dimensions[col].width = w

    put(ws, "A1", "作業記録", size=14, bold=True)
    put(ws, "A2", "作業のたびに1行ずつ記入します。対応時間は0.25時間（15分）単位。区分と状態はプルダウンから選びます。", size=9, color="666666")

    headers = ["No", "日付", "依頼者", "区分", "内容", "対応時間", "状態"]
    for col, h in zip("ABCDEFG", headers):
        put(ws, f"{col}4", h, bold=True, color="FFFFFF", fill=HEAD_FILL, align="center", border=BOX)

    dv_cat = DataValidation(type="list", formula1=f"={S}$F$6:$F${5 + len(CATEGORIES)}", allow_blank=True)
    dv_st = DataValidation(type="list", formula1='"' + ",".join(STATUSES) + '"', allow_blank=True)
    dv_h = DataValidation(type="decimal", operator="between", formula1="0", formula2="24", allow_blank=True,
                          error="0〜24の数値（時間）を入力してください", showErrorMessage=True)
    for dv in (dv_cat, dv_st, dv_h):
        ws.add_data_validation(dv)

    # 記入例（スタンダードプランで月13.5時間 = 請求書テンプレートの記入例と同じ）
    examples = [
        (date(2026, 10, 2), "総務部 A様", "アカウント管理", "新入社員2名のアカウント発行", 1.0, "完了"),
        (date(2026, 10, 5), "営業部 B様", "ヘルプデスク", "共有フォルダにアクセスできない件の調査・修正", 0.5, "完了"),
        (date(2026, 10, 9), "総務部 A様", "PCセットアップ", "新入社員用PC 2台の初期設定", 3.0, "完了"),
        (date(2026, 10, 14), "代表 C様", "セキュリティ", "多要素認証の全社展開（準備・手順書作成）", 2.5, "完了"),
        (date(2026, 10, 19), "経理部 D様", "ヘルプデスク", "複合機のスキャン送信設定", 0.75, "完了"),
        (date(2026, 10, 22), "総務部 A様", "アカウント管理", "退職者1名のアカウント停止・データ移管", 0.5, "完了"),
        (date(2026, 10, 26), "代表 C様", "定例・相談", "月次定例（オンライン）", 1.0, "完了"),
        (date(2026, 10, 28), "営業部 B様", "クラウド運用", "Teams のチーム・チャネル構成の整理", 2.0, "対応中"),
        (date(2026, 10, 30), "代表 C様", "セキュリティ", "多要素認証の全社展開（実施・問い合わせ対応）", 2.25, "完了"),
    ]
    for i, r in enumerate(range(LOG_FIRST, LOG_LAST + 1)):
        ex = examples[i] if i < len(examples) else (None,) * 6
        put(ws, f"A{r}", i + 1, align="center", border=BOX, size=9)
        inp(ws, f"B{r}", ex[0], fmt=MD, border=BOX, align="center")
        inp(ws, f"C{r}", ex[1], border=BOX)
        inp(ws, f"D{r}", ex[2], border=BOX)
        inp(ws, f"E{r}", ex[3], border=BOX, wrap=True)
        inp(ws, f"F{r}", ex[4], fmt=HOURS, border=BOX)
        inp(ws, f"G{r}", ex[5], border=BOX, align="center")
        dv_cat.add(f"D{r}")
        dv_h.add(f"F{r}")
        dv_st.add(f"G{r}")

    r = LOG_LAST + 1
    put(ws, f"E{r}", "合計", bold=True, align="right", border=BOX, fill=SOFT_FILL)
    put(ws, f"F{r}", f"=SUM(F{LOG_FIRST}:F{LOG_LAST})", fmt=HOURS, bold=True, border=BOX, fill=SOFT_FILL)
    return ws


# ---------------------------------------------------------------- 月次レポート
def build_report(wb):
    ws = wb.active
    ws.title = "月次レポート"
    ws.sheet_view.showGridLines = False
    for col, w in {"A": 2, "B": 20, "C": 14, "D": 14, "E": 14, "F": 14, "G": 14}.items():
        ws.column_dimensions[col].width = w

    rng_cat = f"{L}$D${LOG_FIRST}:$D${LOG_LAST}"
    rng_h = f"{L}$F${LOG_FIRST}:$F${LOG_LAST}"
    rng_desc = f"{L}$E${LOG_FIRST}:$E${LOG_LAST}"
    rng_st = f"{L}$G${LOG_FIRST}:$G${LOG_LAST}"

    put(ws, "B1", "月次作業報告書", size=20, bold=True, align="center")
    ws.merge_cells("B1:G1")
    put(ws, "B2", "黄色のセルが入力欄です。数値は「作業記録」シートから自動で集計されます。", size=8, color="999999")

    inp(ws, "B4", "サンプル商事株式会社", size=13, bold=True, border=BOTTOM)
    ws.merge_cells("B4:D4")
    put(ws, "E4", "御中", size=12, border=BOTTOM)

    put(ws, "F3", "対象月", size=9, align="right")
    inp(ws, "G3", date(2026, 10, 1), fmt=MONTH, size=9)
    put(ws, "F4", "報告日", size=9, align="right")
    inp(ws, "G4", date(2026, 11, 5), fmt=DATE, size=9)
    put(ws, "F5", "ラクシス", size=10, bold=True, align="right")
    ws.merge_cells("F5:G5")
    inp(ws, "F6", "〔氏名〕", size=9, align="right")
    ws.merge_cells("F6:G6")

    put(ws, "B6", "ご契約プラン", bold=True, size=9)
    inp(ws, "C6", "スタンダード", size=10, bold=True)
    dv_plan = DataValidation(type="list", formula1=f"={S}$B$6:$B$8", allow_blank=False)
    ws.add_data_validation(dv_plan)
    dv_plan.add("C6")

    # ---- サマリー
    put(ws, "B8", "1. 今月のサマリー", size=12, bold=True)
    labels = ["対応件数", "作業時間", "月間対応時間", "利用率", "超過時間", "超過料金（税別）"]
    formulas_ = [
        f'=COUNTA({rng_desc})',
        f"=SUM({rng_h})",
        f"=INDEX({S}$C$6:$C$8,MATCH($C$6,{S}$B$6:$B$8,0))",
        "=IF(D10=0,0,C10/D10)",
        "=MAX(0,C10-D10)",
        f"=F10*{S}$C$10",
    ]
    fmts = ['0"件"', HOURS, HOURS, "0%", HOURS, YEN]
    for i, (lab, fx, fm) in enumerate(zip(labels, formulas_, fmts)):
        col = "BCDEFG"[i]
        put(ws, f"{col}9", lab, bold=True, size=9, fill=SOFT_FILL, border=BOX, align="center")
        put(ws, f"{col}10", fx, fmt=fm, size=14, bold=True, border=BOX, align="center")
    ws.row_dimensions[10].height = 30
    put(ws, "B11", '=IF(F10>0,"今月は月間対応時間を超えたため、超過分を請求書に含めています。","月間対応時間内で対応しました。")', size=9, color="444444")
    ws.merge_cells("B11:G11")

    # ---- 区分別
    put(ws, "B13", "2. 区分別の内訳", size=12, bold=True)
    for col, lab in zip("BCD", ["区分", "件数", "時間"]):
        put(ws, f"{col}14", lab, bold=True, size=9, fill=SOFT_FILL, border=BOX, align="center")
    for i in range(len(CATEGORIES)):
        r = 15 + i
        put(ws, f"B{r}", f"={S}F{6 + i}", size=9, border=BOX)
        put(ws, f"C{r}", f"=COUNTIF({rng_cat},B{r})", fmt='0"件"', size=9, border=BOX, align="center")
        put(ws, f"D{r}", f"=SUMIF({rng_cat},B{r},{rng_h})", fmt=HOURS, size=9, border=BOX, align="center")
    last_cat = 14 + len(CATEGORIES)
    # 0件の区分は薄く表示
    ws.conditional_formatting.add(
        f"B15:D{last_cat}", FormulaRule(formula=["$C15=0"], font=Font(name=FONT, size=9, color="BBBBBB"))
    )
    # 右側：状態
    put(ws, "F14", "状態", bold=True, size=9, fill=SOFT_FILL, border=BOX, align="center")
    put(ws, "G14", "件数", bold=True, size=9, fill=SOFT_FILL, border=BOX, align="center")
    for i, st in enumerate(STATUSES):
        r = 15 + i
        put(ws, f"F{r}", st, size=9, border=BOX, align="center")
        put(ws, f"G{r}", f"=COUNTIF({rng_st},F{r})", fmt='0"件"', size=9, border=BOX, align="center")

    # ---- 文章欄
    blocks = [
        ("3. 主な作業", [
            "・新入社員2名の受け入れ対応（アカウント発行、PC 2台の初期設定）",
            "・多要素認証を全社に展開（手順書を作成し、全員の設定を完了）",
            "・退職者1名のアカウント停止と、メール・ファイルの引き継ぎ",
        ]),
        ("4. 対応中・持ち越しの案件", [
            "・Teams のチーム・チャネル構成の整理（11月中に完了予定）",
        ]),
        ("5. 来月のご提案", [
            "・PC の OS 更新状況の一斉確認（未更新の端末が2台あります）",
            "・共有フォルダの権限見直し（部署異動に合わせて）",
        ]),
        ("6. お気づきの点・お願い", [
            "・多要素認証の設定後、スマホを機種変更する際は事前にご連絡ください。",
        ]),
    ]
    r = last_cat + 2
    for title, lines in blocks:
        put(ws, f"B{r}", title, size=12, bold=True)
        r += 1
        for i in range(max(3, len(lines))):
            inp(ws, f"B{r}", lines[i] if i < len(lines) else None, size=9)
            ws.merge_cells(f"B{r}:G{r}")
            r += 1
        r += 1

    put(ws, f"B{r}", "※ 作業の明細は別紙「作業記録」をご覧ください。", size=8, color="666666")

    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.orientation = "portrait"
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.print_area = f"A1:G{r}"
    return ws


def main():
    wb = Workbook()
    build_report(wb)
    build_log(wb)
    build_settings(wb)
    wb.calculation.fullCalcOnLoad = True
    out = "out/月次作業報告書.xlsx"
    wb.save(out)
    print(out)


if __name__ == "__main__":
    main()
